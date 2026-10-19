r"""Appendix table: the T-ood blocks of the policies that learned the task,
each against its OWN main test, with the Newcombe interval of the difference.

Two steps, because no one Python here has both libraries:

  .venv-analysis\Scripts\python.exe scripts/tools/table_ood_newcombe.py --csv
      reads daten/endstudie_rl_games/evaluation/summary.json (rates, 95 % Newcombe as eval_stats
      wrote them, never recomputed) and adds the Bonferroni reading over all
      comparisons of the table (statsmodels, the same call as eval_stats);
      writes docs/tabellen/Tab_A_ausserhalb_newcombe.csv
  python scripts/tools/table_ood_newcombe.py
      (system Python, python-docx) turns that CSV into
      docs/tabellen/Tab_A_ausserhalb_newcombe.docx

The learned policies come from eval_figures.data_learned's rule (LEARNED_MIN on
the own main test), read from eval_figures, not restated.
"""
import argparse
import csv
import json
import pathlib
import sys

REPO = pathlib.Path(__file__).resolve().parents[2]
OUT = REPO / "docs" / "tabellen" / "Tab_A_ausserhalb_newcombe"
BLOCKS = (("lat_r_hi", "Startversatz 45 mm"), ("yaw_lo", "Gierung −7,5°"), ("yaw_hi", "Gierung +7,5°"),
          ("tilt_hi", "Kippung 15°"), ("friction_lo", "Reibwert 0,05"), ("friction_hi", "Reibwert 0,23"))
LABEL = {"autodr": "Auto-DR", "fixeddr": "Fixed-DR"}
COLS = ("Block", "Policy", "Erfolge", "Erfolgsrate in %", "Haupttest in %", "Differenz in Prozentpunkten",
        "Newcombe-95-%-Intervall", "ganz unter null", "auch nach Bonferroni")


def de(x: float) -> str:
    return f"{x:.1f}".replace(".", ",").replace("-", "−")


def write_csv(summary: pathlib.Path) -> None:
    from statsmodels.stats.proportion import confint_proportions_2indep
    sys.path.insert(0, str(REPO / "scripts"))
    import eval_figures as EF
    pols = json.loads(summary.read_text(encoding="utf-8"))["OOD"]["policies"]
    learned = [p for p in pols.values() if p["reference"]["rate"] >= EF.LEARNED_MIN]
    learned.sort(key=lambda p: (list(LABEL).index(p["policy"]["method"]), p["policy"]["seed"]))
    m = len(learned) * len(BLOCKS)
    rows = []
    for key, word in BLOCKS:
        for p in learned:
            b, r = p["blocks"][key], p["reference"]
            lo, hi = b["newcombe_ci"]
            _lb, hib = confint_proportions_2indep(b["k"], b["n"], r["k"], r["n"], method="newcomb",
                                                  compare="diff", alpha=0.05 / m)
            rows.append((word, f"{LABEL[p['policy']['method']]} Seed {p['policy']['seed']}", f"{b['k']}/{b['n']}",
                         de(100 * b["rate"]), de(100 * r["rate"]), de(100 * b["diff_vs_final"]),
                         f"[{de(100 * lo)}; {de(100 * hi)}]", "ja" if hi < 0 else "nein", "ja" if hib < 0 else "nein"))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT.with_suffix(".csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(COLS)
        w.writerows(rows)
    print(f"[table_ood_newcombe] {len(learned)} policies x {len(BLOCKS)} blocks = {m} comparisons, "
          f"Bonferroni alpha 0.05/{m}; wrote {OUT.with_suffix('.csv')}")


def write_docx() -> None:
    from docx import Document
    from docx.shared import Pt
    with open(OUT.with_suffix(".csv"), encoding="utf-8") as f:
        rows = list(csv.reader(f))
    doc = Document()
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
            for run in c.paragraphs[0].runs:
                run.font.size = Pt(8)
                run.font.bold = i == 0
    doc.save(OUT.with_suffix(".docx"))
    print(f"[table_ood_newcombe] {len(rows) - 1} rows; wrote {OUT.with_suffix('.docx')}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--csv", action="store_true", help="step 1 (.venv-analysis): summary.json -> CSV")
    ap.add_argument("--summary", default=str(REPO / "daten" / "endstudie_rl_games" / "evaluation" / "summary.json"))
    a = ap.parse_args()
    if a.csv:
        write_csv(pathlib.Path(a.summary))
    else:
        write_docx()
    return 0


if __name__ == "__main__":
    sys.exit(main())
