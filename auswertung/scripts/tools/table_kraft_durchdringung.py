r"""Tables of force and interpenetration over ALL
trials (successes and failures) per training condition, main test beside the
T-ood blocks / the doubled observation errors.

One row per test and condition, over its policies with >= 90 % in their own main
test: range over the policies of the median and of the 95 % value of the
peak filtered force, trials at the 30 N limit (summed over the policies), range of
the median of the largest interpenetration. Every number comes from
``plot_kraft_durchdringung.data_h24`` / ``data_h23`` (nearest rank, counts held
against summary.json there), never recomputed here.

Generate the CSV first, then the Word table (requires python-docx):

  python scripts/tools/table_kraft_durchdringung.py --csv
      -> docs/tabellen/Tab_A_ausserhalb_kraft.csv, Tab_A_beobachtungsfehler_kraft.csv
  python scripts/tools/table_kraft_durchdringung.py
      (python-docx) -> the same names as .docx


"""
import argparse
import csv
import json
import pathlib
import sys

REPO = pathlib.Path(__file__).resolve().parents[2]
TABLES = {"H24": REPO / "docs" / "tabellen" / "Tab_A_ausserhalb_kraft",
          "H23": REPO / "docs" / "tabellen" / "Tab_A_beobachtungsfehler_kraft"}
COLS = ("Test", "Verfahren (Policies)", "Kraft, Median in N", "Kraft, 95-%-Wert in N", "Versuche mit ≥ 30 N",
        "Durchdringung, Median in mm")


def de(x: float, nd: int) -> str:
    return f"{x:.{nd}f}".replace(".", ",")


def span(values, nd: int) -> str:
    """'lo–hi' over the policies, one number when both round alike."""
    lo, hi = de(min(values), nd), de(max(values), nd)
    return lo if lo == hi else f"{lo}–{hi}"


def rows_of(d: dict, K) -> list:
    out = []
    for j, (key, label) in enumerate(zip(d["keys"], d["labels"])):
        for i, (m, a) in enumerate(d["arms"].items()):
            f, ip = a["panels"][K.FORCE], a["panels"]["ip_max_mm"]
            out.append((label.replace("\n", " ") if i == 0 else "",
                        f"{K.EF.METHOD_LABEL[m]} ({len(a['seeds'])})",
                        span(f["median"][:, j], 1), span(f["p95"][:, j], 1),
                        f"{int(a['aborts'][:, j].sum())} von {K.EF._thousands(d['n'][key] * len(a['seeds']))}",
                        span(ip["median"][:, j], 2)))
    return out


def write_csv() -> None:
    import pandas as pd
    sys.path.insert(0, str(REPO / "scripts"))
    import plot_kraft_durchdringung as K
    summary = json.loads((K.RESULTS / "summary.json").read_text(encoding="utf-8"))
    df = pd.read_csv(K.RESULTS / "episodes_all.csv", keep_default_na=False, dtype={"tag": str, "boundary_key": str})
    data = {"H24": K.data_h24(summary, df, K.EF._ood_inputs(K.RESULTS, K.EF._ood_sidecar_path())),
            "H23": K.data_h23(summary, df)}
    for key, path in TABLES.items():
        rows = rows_of(data[key], K)
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path.with_suffix(".csv"), "w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(COLS)
            w.writerows(rows)
        print(f"[table_kraft_durchdringung] {len(rows)} rows; wrote {path.with_suffix('.csv')}")


def write_docx() -> None:
    from docx import Document
    from docx.shared import Pt
    for path in TABLES.values():
        with open(path.with_suffix(".csv"), encoding="utf-8") as fh:
            rows = list(csv.reader(fh))
        doc = Document()
        t = doc.add_table(rows=0, cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            cells = t.add_row().cells
            for c, v in zip(cells, row):
                c.text = v
                for run in c.paragraphs[0].runs:
                    run.font.size = Pt(8)
                    run.font.bold = i == 0
        doc.save(path.with_suffix(".docx"))
        print(f"[table_kraft_durchdringung] {len(rows) - 1} rows; wrote {path.with_suffix('.docx')}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--csv", action="store_true", help="step 1: data -> CSV")
    a = ap.parse_args()
    if a.csv:
        write_csv()
    else:
        write_docx()
    return 0


if __name__ == "__main__":
    sys.exit(main())
